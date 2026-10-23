"""Printable label sheet: one label per fish with a QR code + the same text in large print,
so engine/labels.py reads the QR (no errors) and a person (or the OCR fallback) reads the text.
python -m guppycolor.labelsheet LIST.xlsx|csv|txt OUT.pdf [--date 20/03] [--id-col male.id] [--group-col group]
Print at 100% (no 'fit to page'); cut along the grey lines."""
import argparse, csv, os, re
import numpy as np, cv2
from PIL import Image, ImageDraw, ImageFont

ID_PATTERN = r'[Mm]\d{2,4}'
LABEL_MM = (34, 14)            # w, h: fits beside a guppy on the mm paper; QR 12 mm = ~0.4 mm per module
PAGE_MM = (210, 297)
MARGIN_MM = 10

def _font(px, bold=True):
    names = (['arialbd.ttf', 'DejaVuSans-Bold.ttf'] if bold else ['arial.ttf', 'DejaVuSans.ttf'])
    for n in names + ['/usr/share/fonts/truetype/dejavu/' + names[-1]]:
        try: return ImageFont.truetype(n, px)
        except OSError: pass
    return ImageFont.load_default(px)

def qr_payload(row):
    """'M116 G4 20/03': plain text, so any phone app shows it too; read_labels splits it with its patterns."""
    return ' '.join(str(row[k]) for k in ('id', 'group', 'date') if row.get(k))

def label_image(row, dpi=600, mm=LABEL_MM):
    px = dpi/25.4
    W, H = int(mm[0]*px), int(mm[1]*px)
    im = Image.new('L', (W, H), 255); d = ImageDraw.Draw(im)
    q = cv2.QRCodeEncoder.create().encode(qr_payload(row))       # 1 px per module, with quiet zone
    side = int((mm[1] - 2)*px)
    qi = Image.fromarray(q).resize((side, side), Image.NEAREST)
    im.paste(qi, (int(1*px), (H - side)//2))
    x = int(1*px) + side + int(1.5*px)
    d.text((x, int(1.2*px)), str(row['id']), font=_font(int(5.5*px)), fill=0)
    small = '  '.join(str(row[k]) for k in ('group', 'date') if row.get(k))
    if small: d.text((x, int(8.3*px)), small, font=_font(int(3.6*px)), fill=0)
    return im

def make_sheet(rows, out_pdf, dpi=600, mm=LABEL_MM):
    """rows: [{'id': 'M116', 'group': 'G4', 'date': '20/03'}, ...] -> multi-page A4 PDF. Returns the page count."""
    px = dpi/25.4
    PW, PH = int(PAGE_MM[0]*px), int(PAGE_MM[1]*px)
    lw, lh = int(mm[0]*px), int(mm[1]*px); gap = int(2*px); m = int(MARGIN_MM*px)
    cols = (PW - 2*m + gap)//(lw + gap); nrow = (PH - 2*m + gap)//(lh + gap)
    pages = []
    for p0 in range(0, len(rows), cols*nrow):
        page = Image.new('L', (PW, PH), 255); d = ImageDraw.Draw(page)
        for k, row in enumerate(rows[p0:p0 + cols*nrow]):
            x = m + (k % cols)*(lw + gap); y = m + (k//cols)*(lh + gap)
            page.paste(label_image(row, dpi, mm), (x, y))
            d.rectangle([x - 1, y - 1, x + lw, y + lh], outline=170)       # cutting guide
        pages.append(page)
    pages[0].save(out_pdf, 'PDF', resolution=dpi, save_all=True, append_images=pages[1:])
    return len(pages)

def read_rows(path, id_col=None, group_col=None, date=None, id_pattern=ID_PATTERN):
    """Fish list from xlsx/csv/txt. With id_col: that column (and group_col); otherwise every cell or word
    matching id_pattern, once each, in file order."""
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        table = [[('' if v is None else str(v)) for v in r] for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
    elif ext == '.csv':
        text = open(path, newline='', encoding='utf-8-sig').read()
        table = list(csv.reader(text.splitlines(), delimiter=';' if text.count(';') > text.count(',') else ','))
    else:
        table = [[w] for w in open(path, encoding='utf-8').read().split()]
    rows, seen = [], set()
    head = None
    for r in table:
        if id_col and id_col in r:             # header row (each Excel sheet has its own)
            head = r; continue
        if id_col:
            if head is not None and len(r) > head.index(id_col):
                v = r[head.index(id_col)].strip()
                if re.fullmatch(id_pattern, v):
                    g = r[head.index(group_col)].strip() if group_col and group_col in head and len(r) > head.index(group_col) else ''
                    _add(rows, seen, v, g, date)
            continue
        for cell in r:
            for v in re.findall(r'\b' + id_pattern + r'\b', cell): _add(rows, seen, v, '', date)
    return rows

def _add(rows, seen, v, g, date):
    v = v.upper()
    if v in seen: return
    seen.add(v)
    if g and re.fullmatch(r'\d+', g): g = 'G' + g
    rows.append({'id': v, 'group': g, 'date': date or ''})

def main(argv=None):
    ap = argparse.ArgumentParser(prog='guppycolor.labelsheet', description='foglio di etichette QR + testo')
    ap.add_argument('lista'); ap.add_argument('pdf')
    ap.add_argument('--date', default=''); ap.add_argument('--id-col'); ap.add_argument('--group-col')
    a = ap.parse_args(argv)
    rows = read_rows(a.lista, a.id_col, a.group_col, a.date)
    n = make_sheet(rows, a.pdf)
    print(f'{len(rows)} etichette, {n} pagine: {a.pdf}')

if __name__ == '__main__':
    main()
