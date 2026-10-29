"""Analysis of one photo: colour correction -> fish mask -> scale -> measures -> quality flags.
Results are stored in pixels; conversion to mm happens at export, so a doubtful scale can fall back
to the median of its session without re-analysing."""
import csv
import os
import statistics
import time

import cv2
import numpy as np
from PIL import Image, ImageOps

from . import metadata, profile, sizes
from .engine import colour, morpho, scale

CHART_MAX_ERR = 0.1      # find_chart pattern error above this = no usable chart
ASPECT_RANGE = (3, 10)   # fish length / max depth outside this range = suspicious mask
VERSION = 1


def load_image(path):
    with Image.open(path) as im:
        return np.array(ImageOps.exif_transpose(im).convert('RGB'))


# --- colour reference (same CSV layout as the ImageJ plugin: Patch,Name,R,G,B in 0-255) ---
def patch_names():
    cc = colour.colour.CCS_COLOURCHECKERS['ColorChecker24 - After November 2014']
    return list(cc.data.keys())


def save_reference(path, meas):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['Patch', 'Name', 'R', 'G', 'B'])
        for i, (name, rgb) in enumerate(zip(patch_names(), meas.reshape(-1, 3))):
            w.writerow([i + 1, name] + [round(float(v) * 255, 2) for v in rgb])


def load_reference(path):
    if not os.path.exists(path):
        return None
    with open(path, encoding='utf-8') as f:
        rows = list(csv.DictReader(f))
    if len(rows) != 24:
        return None
    return (np.array([[float(r['R']), float(r['G']), float(r['B'])] for r in rows]) / 255).reshape(4, 6, 3)


def set_reference_photo(project, photo):
    """Measure the chart on the chosen reference photo and store it for the whole project."""
    chart = colour.find_chart(load_image(photo))
    if chart is None or chart[0] > CHART_MAX_ERR:
        return False
    save_reference(project.reference_csv, chart[1])
    project.settings['colour']['reference_photo'] = os.path.abspath(photo)
    project.save()
    return True


def loo_error(meas, target):
    """Mean dE00 of each patch predicted by a fit that did not see it (honest correction error)."""
    from scipy.interpolate import RBFInterpolator
    m, t = meas.reshape(-1, 3), target.reshape(-1, 3)
    errs = []
    for k in range(len(m)):
        idx = [j for j in range(len(m)) if j != k]
        f = RBFInterpolator(m[idx], t[idx], kernel='thin_plate_spline', smoothing=1e-3)
        errs.append(colour.delta_e(np.clip(f(m[k:k + 1]), 0, 1)[0], t[k]))
    return float(np.mean(errs))


def correct_colour(img, settings, reference):
    """Returns (image to measure colours on, info dict, flags)."""
    mode = settings['colour']['mode']
    info = {'colour_chart': False, 'correction': 'none', 'chart_error_de00': None}
    flags = []
    if mode == 'none':
        return img, info, flags
    if mode in ('reference', 'shared') and reference is None:
        flags.append('reference_missing')
        return img, info, flags
    if mode == 'shared':          # one transform for every photo, measured once on the reference photo
        info['correction'] = 'shared'
        return colour.apply_tps(colour.fit_tps(reference), img), info, flags
    chart = colour.find_chart(img)
    if chart is None or chart[0] > CHART_MAX_ERR:
        flags.append('no_chart')
        return img, info, flags
    meas = chart[1]
    target = colour.REF if mode == 'nominal' else reference
    info.update(colour_chart=True, correction='nominal' if mode == 'nominal' else 'reference',
                chart_error_de00=round(loo_error(meas, target), 2))
    return colour.apply_tps(colour.fit_tps(meas, None if mode == 'nominal' else reference), img), info, flags


SESSION_FLAGS = ('scale_off_session', 'size_off_session')
SERIOUS = {'label_uncertain', 'fish_at_border', 'mask_in_pieces', 'unusual_shape', 'scale_doubtful', 'scale_not_found',
           'scale_to_set', 'reference_missing', *SESSION_FLAGS}       # these send the photo to Review


def find_scale(img, mask, settings):
    mode = settings['scale']['mode']
    if mode == 'fixed':
        v = settings['scale'].get('px_per_mm')
        return {'px_per_mm': float(v) if v else None, 'scale_source': 'fixed', 'scale_ok': bool(v)}, []
    if mode == 'none':
        return {'px_per_mm': None, 'scale_source': 'none', 'scale_ok': False}, []
    if mode == 'manual':                      # a line drawn in Review (scaleline) gives it
        return {'px_per_mm': None, 'scale_source': 'to_set', 'scale_ok': False}, ['scale_to_set']
    try:
        g = scale.grid_scale(img, mask)
    except Exception:
        g = None
    if g is None:
        return {'px_per_mm': None, 'scale_source': 'not_found', 'scale_ok': False}, ['scale_not_found']
    flags = [] if g['ok'] else ['scale_doubtful']
    line = g.get('line')
    return {'px_per_mm': round(g['px_per_mm'], 4), 'scale_source': 'graph_paper', 'scale_ok': g['ok'],
            'scale_squares': g.get('squares'),
            'scale_line': [[round(v, 1) for v in pt] for pt in line] if line else None}, flags


def photo_scale(project, folder, photo, img, mask):
    """Line drawn on this photo > automatic/fixed scale > line drawn for its folder (when the first is not good)."""
    from . import scaleline
    if project.settings['scale']['mode'] != 'none':
        ppm = scaleline.photo_scale(project, photo)
        if ppm:
            return {'px_per_mm': round(ppm, 4), 'scale_source': 'line', 'scale_ok': True}, []
    sc, f = find_scale(img, mask, project.settings)
    if not sc['scale_ok'] and project.settings['scale']['mode'] != 'none':
        ppm = scaleline.folder_scale(project, os.path.basename(folder))
        if ppm:
            return dict(sc, px_per_mm=round(ppm, 4), scale_source='line_folder', scale_ok=True), []
    return sc, f


def load_expected_ids(path, pattern):
    """Valid ids from an Excel, CSV or text file: every cell/word that matches the id pattern."""
    import re
    if not path or not os.path.exists(path):
        return None
    rx = re.compile(pattern, re.I)
    values = []
    if path.lower().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                values += [str(v) for v in row if v is not None]
    else:
        with open(path, encoding='utf-8', errors='replace') as f:
            values = re.split(r'[\s,;]+', f.read())
    ids = sorted({v.strip().upper() for v in values if rx.fullmatch(v.strip())})
    return ids or None


def reads_label(settings):
    """The label in the photo is read when it gives the id, or when asked for its date/group."""
    lb = settings['labels']
    return lb.get('id_source') in ('label', 'table') or bool(lb.get('read_label'))


def date_mismatch(label_date, date_taken):
    """'20/03' on the label vs EXIF '2024-03-21 10:00:00': True when day or month differ."""
    import re
    m = re.match(r'\s*(\d{1,2})\D(\d{1,2})', label_date or '')
    e = re.match(r'(\d{4})-(\d{2})-(\d{2})', date_taken or '')
    return bool(m and e) and (int(m.group(1)), int(m.group(2))) != (int(e.group(3)), int(e.group(2)))


def landmarks_tl(points):
    if 'snout' in points and 'tail_tip' in points:
        return float(np.hypot(*np.subtract(points['snout'], points['tail_tip'])))
    return None


def find_landmarks(project, photo, img, mask, r):
    """Landmarks clicked by hand (screen 8) win; otherwise the landmark model, if one is chosen (sizes from them
    in sizes.update)."""
    from . import landmarks
    pts, src = landmarks.load(project, photo), 'manual'
    name = project.settings.get('landmarks', {}).get('model')
    if not pts and name:
        try:
            pts, src = landmarks.predict(project, name, img, mask), 'model'
        except Exception:
            return ['landmark_model_error']
    if not pts:
        return []
    r.update(landmarks=pts, landmarks_source=src)
    return []


def analysis_mode(r, photo, settings):
    """'size' or 'size_colour' for one photo (see the 'analysis' settings)."""
    import re
    an = settings.get('analysis', {})
    folder = os.path.normcase(os.path.abspath(os.path.dirname(photo)))
    for f in an.get('size_only_folders', []):
        f = os.path.normcase(os.path.abspath(f))
        if folder == f or folder.startswith(f + os.sep):
            return 'size'
    pat = an.get('size_only_pattern')
    if pat:
        for text in (r.get('label_id'), os.path.splitext(r.get('file', ''))[0]):
            try:
                if text and re.search(pat, text, re.I):
                    return 'size'
            except re.error:
                break
    return an.get('default', 'size_colour')


def label_fields(res):
    return {'label_text': res.get('text_raw') or '', 'label_id': res.get('male_id'),
            'label_group': res.get('group'), 'label_date': res.get('date'),
            'label_confidence': round(float(res.get('confidence') or 0), 3), 'label_source': res.get('source'),
            'label_needs_review': bool(res.get('needs_review')),
            'label_candidates': [str(t) for t, _ in (res.get('id_candidates') or [])][:4],
            'label_box': list(res['label_box']) if res.get('label_box') else None}


def label_patterns(settings):
    lb = settings['labels']
    return {k: lb[k] for k in ('male_id', 'group', 'date') if lb.get(k)}


def read_label(photo, settings, expected_ids):
    """Fields of the label written in the photo. Errors never stop the analysis."""
    from .engine import labels
    return label_fields(labels.read_labels(photo, expected_ids=expected_ids, patterns=label_patterns(settings)))


def labels_from_table(project, todo, expected_ids):
    """With a data table (photo name -> fish id): ids from the table, the other photos of the same folder
    from their neighbours in shooting order + the label. Returns {photo: label fields}; {} on failure."""
    from .engine import labels
    st = project.settings
    try:
        table = labels.load_id_table(st['labels']['id_table_file'], st['labels']['male_id'])
    except Exception:
        return {}
    out = {}
    for d in sorted({os.path.dirname(p) for _, p in todo}):
        photos = [p for _, p in project.photos() if os.path.dirname(p) == d]
        try:
            res = labels.fill_from_sequence(photos, table, expected_ids, label_patterns(st))
        except Exception:
            continue
        out.update({p: label_fields(r) for p, r in res.items()})
    return out


def touches_border(mask, margin=3):
    return bool(mask[:margin].any() or mask[-margin:].any() or mask[:, :margin].any() or mask[:, -margin:].any())


# colours of the measures drawn on check/<photo>.jpg (RGB); the Results and Review screens explain them
MEASURE_COLOURS = {'outline': (230, 0, 200), 'midline': (0, 220, 255), 'snout': (0, 200, 0), 'end': (255, 30, 30),
                   'peduncle': (255, 140, 0), 'depth': (255, 230, 0), 'scale': (255, 255, 255)}


def save_overlay(path, img, mask, morph, max_side=1200, margin=0.3, ppm=None):
    """Crop around the fish (it can be ~3% of a wide shot) with the measures drawn: outline (magenta), midline =
    mask length (cyan) from the snout (green) to the end of the outline (red), end of the body = caudal peduncle when
    the fin is in the outline (orange), maximum depth (yellow), 5 mm bar (white).
    Returns the crop box (x0, y0, x1, y1) in original pixels, to map clicks back."""
    ys, xs = np.nonzero(mask)
    pad = int(margin * max(np.ptp(xs), np.ptp(ys)))
    x0, y0 = max(0, xs.min() - pad), max(0, ys.min() - pad)
    x1, y1 = min(img.shape[1], xs.max() + pad + 1), min(img.shape[0], ys.max() + pad + 1)
    vis = img[y0:y1, x0:x1].copy()
    s = max_side / max(vis.shape[:2])
    vis = cv2.resize(vis, None, fx=s, fy=s, interpolation=cv2.INTER_AREA)
    m = cv2.resize(mask[y0:y1, x0:x1].astype(np.uint8), (vis.shape[1], vis.shape[0]), interpolation=cv2.INTER_NEAREST)
    cnts = cv2.findContours(m, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_NONE)[0]
    t = max(2, max(vis.shape[:2]) // 400)
    C = MEASURE_COLOURS
    cv2.drawContours(vis, cnts, -1, C['outline'], t)
    to = (lambda p: (int(round((p[0] - x0) * s)), int(round((p[1] - y0) * s))))
    if morph is not None:
        line = morph.get('line')
        if line is not None:
            cv2.polylines(vis, [np.array([to(p) for p in line], np.int32)], False, C['midline'], t, cv2.LINE_AA)
        half = 0.65 * morph['max_depth_px'] * np.asarray(morph['normal'])
        ticks = [(morph['snout'], 'snout'), (morph['end'], 'end')]
        if morph.get('fin_found'):
            ticks.append((morph['ped'], 'peduncle'))
        for p, k in ticks:
            p = np.asarray(p)
            cv2.line(vis, to(p - half), to(p + half), C[k], t + 1, cv2.LINE_AA)
        if morph.get('depth_line') is not None:
            a, b = morph['depth_line']
            cv2.line(vis, to(a), to(b), C['depth'], t + 1, cv2.LINE_AA)
    if ppm:
        bar = int(round(5 * ppm * s))
        if 0 < bar < vis.shape[1] - 40:
            h = vis.shape[0]
            cv2.line(vis, (20, h - 20), (20 + bar, h - 20), (0, 0, 0), t + 4)
            cv2.line(vis, (20, h - 20), (20 + bar, h - 20), C['scale'], t + 1)
            cv2.putText(vis, '5 mm', (20, h - 20 - 4 * t), cv2.FONT_HERSHEY_SIMPLEX, 0.5 + t / 6, (0, 0, 0), t + 2,
                        cv2.LINE_AA)
            cv2.putText(vis, '5 mm', (20, h - 20 - 4 * t), cv2.FONT_HERSHEY_SIMPLEX, 0.5 + t / 6, C['scale'], t,
                        cv2.LINE_AA)
    Image.fromarray(vis).save(path, quality=88)
    return [int(x0), int(y0), int(x1), int(y1)]


def mask_flags(mask, morph):
    """Quality checks on the mask itself (the detector's confidence is low even on good masks)."""
    flags = []
    if touches_border(mask):
        flags.append('fish_at_border')
    n, _, st, _ = cv2.connectedComponentsWithStats(mask.astype(np.uint8), 8)
    areas = sorted(st[1:, cv2.CC_STAT_AREA], reverse=True)
    if len(areas) > 1 and areas[1] > 0.2 * areas[0]:
        flags.append('mask_in_pieces')
    ratio = morph['tl_px'] / max(morph['max_depth_px'], 1e-9)
    if not ASPECT_RANGE[0] <= ratio <= ASPECT_RANGE[1]:
        flags.append('unusual_shape')
    return flags


def image_quality(img, mask, erode=15):
    """On the fish (eroded, away from the outline): % of burnt pixels, sharpness (variance of the Laplacian),
    median brightness. Sharpness depends on resolution, so it is judged against the session median at export."""
    inner = cv2.erode(mask.astype(np.uint8), np.ones((2 * erode + 1,) * 2, np.uint8)) > 0
    if inner.sum() < 100:
        inner = mask
    mx = img.max(-1)
    lap = cv2.Laplacian(cv2.cvtColor(img, cv2.COLOR_RGB2GRAY), cv2.CV_32F)
    return {'saturated_pct': round(100 * float((mx[inner] >= 250).mean()), 2),
            'sharpness': round(float(lap[inner].var()), 2),
            'fish_brightness': round(float(np.median(mx[inner])), 1)}


def manual_mask_path(project, photo):
    return os.path.join(project.dir('masks'), project.photo_id(photo) + '_manual.png')


def save_manual_mask(project, photo, mask):
    """A mask fixed by the user (e.g. one click in Revisione): used instead of the detector from now on."""
    cv2.imwrite(manual_mask_path(project, photo), mask.astype(np.uint8) * 255)


COLOUR_FLAGS = ('colours_uncertain', 'colour_start_from_snout', 'colour_start_missing')


def measure_colours(project, r, corrected, mask):
    """Colour classes with the project's trained model; areas stay in px (mm at export). The head before the
    profile's colour start (e.g. 2 mm behind the eye) counts as fish without colour."""
    from . import training
    cs = project.settings['colours']
    prof = profile.load(project.settings['profile'])
    forest = training.load_model(project, cs['model'])
    flags = []
    start = sizes.colour_start(r, r.get('landmarks'), prof, r.get('px_per_mm'))
    head = None
    r.pop('colour_start_line', None)
    if start:
        point, u, src = start
        head = sizes.ahead_of(mask.shape, point, u)
        n = (-u[1], u[0])
        half = 0.5 * float(r.get('max_depth_px') or 100)
        r['colour_start_line'] = [[round(point[0] - half * n[0], 1), round(point[1] - half * n[1], 1)],
                                  [round(point[0] + half * n[0], 1), round(point[1] + half * n[1], 1)]]
        if src == 'snout':
            flags.append('colour_start_from_snout')
    elif prof.get('colour_start'):
        flags.append('colour_start_missing')
    r['colour_start_source'] = start[2] if start else None
    colori, labmap, conf = training.measure_colours(corrected, mask, forest, r.get('px_per_mm') if r.get('scale_ok') else None,
                                                   cs.get('merge_mm', 0), cs.get('min_spot_mm2', 0),
                                                   training.classes_off(project), head=head)
    r.update(colours=colori, colour_model=cs['model'], colour_confidence=round(conf, 3))
    try:                            # class map kept for the pattern screen, if installed
        from . import pattern
        r['colour_map'] = pattern.save_colour_map(project, r['photo_id'], labmap)
    except ImportError:
        pass
    training.colour_overlay(os.path.join(project.dir('check'), r['photo_id'] + '_colours.jpg'), corrected, labmap, forest)
    return flags + (['colours_uncertain'] if conf < 0.6 else [])


def outline_problems(mask):
    """Quality flags of an outline alone (border, pieces, shape), to choose between two outlines."""
    try:
        return mask_flags(mask, morpho.measure(mask))
    except Exception:
        return ['unusual_shape']


def find_outline(project, img, segmenter=None):
    """Fish mask from the imported outline model when one is in use (with its parts, e.g. body and caudal fin),
    else from the zero-training detector (YOLOE + SAM). Returns (mask, parts, confidence, source, flags); a model
    that finds no fish or fails falls back to the detector, with a flag."""
    from .engine import segment
    name = project.settings.get('outline', {}).get('model')
    flags = []
    if name and segmenter is None:
        from . import outline
        try:
            parts, conf = outline.predict(project, name, img)
            if parts:
                mask = np.logical_or.reduce(list(parts.values()))
                bad = outline_problems(mask)
                if not bad:
                    return mask, parts, conf, name, flags
                other, _, conf2 = (segmenter or segment.segment_fish)(img)    # the detector may do better
                if other is None or not other.any() or len(outline_problems(other)) >= len(bad):
                    return mask, parts, conf, name, flags
                return other, {}, conf2, 'detector', flags + ['outline_model_rejected']
            else:
                flags.append('outline_model_missed')
        except Exception:
            flags.append('outline_model_error')
    mask, _, conf = (segmenter or segment.segment_fish)(img)
    return mask, {}, conf, 'detector', flags


def guide_outline(img, mask, points, grow=None, tolerance=0.02):
    """Landmarks that fall outside the outline (e.g. a tail tip on a transparent caudal fin SAM left out) are used as
    clicks to add the missing part, as "Add to the outline" does in Review. A landmark closer to the outline than
    `tolerance` of the fish length counts as inside. Returns (mask, keys of the landmarks that grew it)."""
    if not points or mask is None or not mask.any():
        return mask, []
    if grow is None:
        from .engine import segment
        grow = segment.add_to_mask
    ys, xs = np.nonzero(mask)
    near = max(3.0, tolerance * float(np.hypot(np.ptp(xs), np.ptp(ys))))
    used = []
    for k, (x, y) in points.items():
        xi, yi = int(round(x)), int(round(y))
        if not (0 <= yi < mask.shape[0] and 0 <= xi < mask.shape[1]) or mask[yi, xi]:
            continue
        d = cv2.distanceTransform((~mask).astype(np.uint8), cv2.DIST_L2, 5)
        if d[yi, xi] <= near:
            continue
        try:
            new = grow(img, mask, xi, yi)
        except Exception:
            new = None
        if new is not None and new[yi, xi] and not touches_border(new) and new.sum() < 2 * mask.sum():
            mask, used = new, used + [k]
    return mask, used


def analyse(project, folder, photo, segmenter=None, reference=None, expected_ids=None, label=None):
    """Full analysis of one photo; returns the result dict (also saved in the project)."""
    t0 = time.time()
    st = project.settings
    pid = project.photo_id(photo)
    old = project.load_result(photo) or {}
    r = {'version': VERSION, 'photo_id': pid, 'file': os.path.basename(photo), 'path': os.path.abspath(photo),
         'folder': os.path.basename(folder), 'settings_hash': project.settings_hash(),
         'source_mtime': os.path.getmtime(photo)}
    r.update(metadata.parse_filename(photo, st.get('filename_pattern', '')))
    r.update(metadata.read_exif(photo))
    flags = []
    try:
        img = load_image(photo)
        r['width_px'], r['height_px'] = int(img.shape[1]), int(img.shape[0])
        lb = st['labels']
        if lb.get('id_source') == 'filename' and r.get('id'):
            r.update(label_id=r['id'], label_source='filename')
        if reads_label(st):
            if old.get('label_manual'):              # id typed by the user wins
                r.update({k: v for k, v in old.items() if k.startswith('label_')})
            else:
                if expected_ids is None:
                    expected_ids = load_expected_ids(lb.get('expected_ids_file'), lb['male_id'])
                try:
                    res = label if label is not None else read_label(photo, st, expected_ids)
                    if lb.get('id_source') == 'filename' and r.get('label_id'):      # the file name gives the id
                        res = {k: v for k, v in res.items() if k not in ('label_id', 'label_source', 'label_needs_review')}
                    r.update(res)
                except Exception as e:
                    r['label_error'] = f'{type(e).__name__}: {e}'[:200]
                    r['label_needs_review'] = True
                if r.get('label_needs_review'):
                    flags.append('label_uncertain')
            if date_mismatch(r.get('label_date'), r.get('date_taken')):
                flags.append('date_mismatch')
        if reference is None and st['colour']['mode'] in ('reference', 'shared'):
            reference = load_reference(project.reference_csv)
        corrected, info, f = correct_colour(img, st, reference)
        r.update(info); flags += f
        manual = manual_mask_path(project, photo)
        parts = {}
        if os.path.exists(manual):
            mask, conf = cv2.imread(manual, cv2.IMREAD_GRAYSCALE) > 0, 1.0
            r.update(manual_mask=True, outline_source='manual')
        else:
            mask, parts, conf, src, f = find_outline(project, img, segmenter)
            r['outline_source'] = src; flags += f
        r['fish_confidence'] = round(float(conf), 3)
        if mask is None or not mask.any():
            r.update(status='error', flags=';'.join(flags + ['fish_not_found']), seconds=round(time.time() - t0, 1))
            project.save_result(photo, r)
            return r
        flags += find_landmarks(project, photo, img, mask, r)
        if not r.get('manual_mask') and segmenter is None:     # landmarks outside the outline add the missing part
            mask, grown = guide_outline(img, mask, r.get('landmarks'))
            if grown:
                r['outline_grown_to'] = grown
                flags.append('outline_grown_to_landmarks')
        sc, f = photo_scale(project, folder, photo, img, mask)
        r.update(sc); flags += f
        morph = morpho.measure(mask, img, head=None if st.get('head', 'left') == 'auto' else st.get('head', 'left'))
        r.update(fish_area_px=morph['area_px'], mask_length_px=round(morph['tl_px'], 1))
        if 'body' in parts and len(parts) > 1:        # the outline model drew the parts: they win over the guess
            part_keys = {pt['key'] for pt in profile.load(st['profile']).get('parts', [])}
            morph['body_area_px'] = int(parts['body'].sum())
            morph['fin_found'] = bool(parts.get('caudal') is not None and parts['caudal'].any())
            for k, m in parts.items():
                if k != 'body' and m.any() and k in part_keys and profile.SAFE_NAME.fullmatch(k):    # from its json
                    cv2.imwrite(os.path.join(project.dir('masks'), f'{pid}_{k}.png'), m.astype(np.uint8) * 255)
        r['outline'] = {'body_area_px': morph['body_area_px'], 'max_depth_px': round(morph['max_depth_px'], 1),
                        'mask_length_px': r['mask_length_px'], 'fin_found': morph['fin_found'],
                        'axis': [round(float(v), 4) for v in morph['axis']],
                        'snout': [round(float(v), 1) for v in morph['snout']],
                        'midline': np.round(np.asarray(morph['line'])[::5], 1).tolist()}
        flags = sizes.update(r, mask, r.get('landmarks'), flags)     # landmarks win over the outline
        flags += mask_flags(mask, morph)
        r.update(image_quality(img, mask))
        if r['saturated_pct'] > 1:
            flags.append('overexposed')
        cv2.imwrite(os.path.join(project.dir('masks'), pid + '.png'), mask.astype(np.uint8) * 255)
        r['check_box'] = save_overlay(os.path.join(project.dir('check'), pid + '.jpg'), corrected, mask, morph,
                                      ppm=r.get('px_per_mm') if r.get('scale_ok') else None)
        r['analysis'] = old.get('analysis_manual') or analysis_mode(r, photo, st)
        if old.get('analysis_manual'):
            r['analysis_manual'] = old['analysis_manual']
        if st['colours'].get('model') and r['analysis'] != 'size':
            flags += measure_colours(project, r, corrected, mask)
        r['status'] = 'to_check' if SERIOUS & set(flags) else 'ok'
    except Exception as e:      # one broken photo must not stop the batch
        r['status'] = 'error'
        flags.append('error:' + type(e).__name__)
        r['error'] = str(e)[:300]
    if old.get('review') in ('reviewed', 'excluded'):      # keep the user's review decision
        r['review'] = old['review']
        r['status'] = old['review']
    r['flags'] = ';'.join(flags)
    r['seconds'] = round(time.time() - t0, 1)
    project.save_result(photo, r)
    return r


def run(project, segmenter=None, progress=None, stop=None):
    """Analyse every pending photo. progress(i, n, result) after each; stop() -> True interrupts."""
    todo = project.pending()
    if todo and segmenter is None:
        from .engine import segment
        segment.models()            # download/verify the weights once; a failure stops the run instead of failing every photo
    reference = load_reference(project.reference_csv)
    lb = project.settings['labels']
    expected = load_expected_ids(lb.get('expected_ids_file'), lb['male_id']) if reads_label(project.settings) else None
    from_table = labels_from_table(project, todo, expected) if lb.get('id_source') == 'table' and lb.get('id_table_file') else {}
    for i, (folder, photo) in enumerate(todo):
        if stop is not None and stop():
            break
        r = analyse(project, folder, photo, segmenter=segmenter, reference=reference, expected_ids=expected,
                    label=from_table.get(photo))
        if progress is not None:
            progress(i + 1, len(todo), r)
    session_checks(project)
    return len(todo)


def session_checks(project, results=None):
    """Flag the photos whose scale (px/mm) or fish length (SL, else outline length, in mm) is more than
    checks.session_pct % away from the median of their session (folder), so a wrong scale shows in Review. Needs 3
    photos in the session. Updates and saves only the results that change; returns the results."""
    pct = float((project.settings.get('checks') or {}).get('session_pct') or 0) / 100
    results = project.results() if results is None else results
    by = {}
    for r in results:
        if r.get('status') != 'error' and r.get('review') != 'excluded' and r.get('px_per_mm'):
            by.setdefault(r['folder'], []).append(r)
    length = lambda r: (r.get('standard_length_px') or r.get('mask_length_px') or 0) / r['px_per_mm']
    for r in results:
        old = [f for f in (r.get('flags') or '').split(';') if f]
        flags = [f for f in old if f not in SESSION_FLAGS]
        group = by.get(r.get('folder'), [])
        if pct > 0 and r in group and len(group) >= 3:
            ppm = statistics.median(o['px_per_mm'] for o in group)
            if abs(r['px_per_mm'] / ppm - 1) > pct:
                flags.append('scale_off_session')
            lens = [length(o) for o in group if length(o) > 0]
            if lens and length(r) > 0 and abs(length(r) / statistics.median(lens) - 1) > pct:
                flags.append('size_off_session')
        if flags != old:
            r['flags'] = ';'.join(flags)
            if not r.get('review') and r.get('status') in ('ok', 'to_check'):
                r['status'] = 'to_check' if SERIOUS & set(flags) else 'ok'
            project.save_result(r['path'], r)
    return results
